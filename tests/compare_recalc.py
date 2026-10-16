# Compare LibreOffice-recalculated regenerated workbooks with the stakeholder's cached values (per staff, per RM column group).
import openpyxl, sys
def load(path):
    wb = openpyxl.load_workbook(path, data_only=True); out = {}
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible': continue
        hdr = None
        for r in range(1, 20):
            for c in range(1, 40):
                if str(ws.cell(r, c).value or '').strip() == 'Employee No': hdr = (r, c); break
            if hdr: break
        if not hdr: continue
        H = {}; npts = 0
        for c in range(1, 45):
            h = ' '.join(str(ws.cell(hdr[0], c).value or '').lower().split())
            if not h: continue
            if h == 'no. of points': H['M' if npts == 0 else 'pts'] = c; npts += 1
            elif h.startswith('pool'): H['pool'] = c
            elif h.startswith('point entitlement'): H['K'] = c
            elif h.startswith('days worked'): H['L'] = c
            elif h.startswith('mia') or h.startswith('deduction') or '0.75' in h: H.setdefault('rb', []).append(c)
            elif h.startswith(('operational', 'permenant', 'casual')) or 'foreign worker' in h or h.startswith('agent'): H.setdefault('ent', []).append(c)
        rows = {}
        for r in range(hdr[0] + 1, hdr[0] + 80):
            name = ws.cell(r, hdr[1] + 1).value
            if any('TOTAL POINTS' in str(ws.cell(r, c).value or '') for c in range(1, 15)): break
            if not name or not isinstance(name, str) or name.startswith('('): continue
            g = lambda c: float(ws.cell(r, c).value or 0) if isinstance(ws.cell(r, c).value, (int, float)) else 0
            rows[name.strip()] = dict(K=g(H['K']), L=g(H['L']), M=g(H['M']), ent=sum(g(c) for c in H['ent']), rb=sum(g(c) for c in H['rb']), pool=g(H['pool']))
        pv = None
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and ('point (fixed)' in cell.value.lower()):
                    for c in range(cell.column + 1, cell.column + 10):
                        v = ws.cell(cell.row, c).value
                        if isinstance(v, (int, float)): pv = v; break
        out[ws.title] = (rows, pv)
    return out
for a, b in [('fixtures/Service Charge - Aug26.xlsx', 'out/sc_recalc.xlsx'), ('fixtures/Incentive - Aug26.xlsx', 'out/inc_recalc.xlsx')]:
    A, B = load(a), load(b); worst = 0; n = 0
    for sheet, (ra, pva) in A.items():
        rb_, pvb = B.get(sheet, ({}, None))
        print(f'{sheet:22s} PV stakeholder={pva} regenerated={pvb}  staff {len(ra)}/{len(rb_)}')
        for name, va in ra.items():
            vb = rb_.get(name)
            if not vb: print('   missing', name); continue
            for k in va:
                d = abs(va[k] - vb[k]); worst = max(worst, d); n += 1
                if d > 0.011: print(f'   {name[:30]:30s} {k}: {va[k]} vs {vb[k]}')
    print(f'{a}: compared {n} values, max abs difference {worst:.4f}\n')
